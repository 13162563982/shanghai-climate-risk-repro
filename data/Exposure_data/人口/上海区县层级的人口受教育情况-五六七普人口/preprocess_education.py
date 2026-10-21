# preprocess_education.py
# 预处理五六七普人口受教育情况数据 —— 处理合并单元格、多行表头
# 输入: 【五普】人口受教育情况.xlsx, 【六普】人口受教育情况.xlsx, 【七普】人口受教育情况.xlsx
# 输出: education_5_cleaned.csv, education_6_cleaned.csv, education_7_cleaned.csv

import os
import openpyxl
import pandas as pd

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SRC_DIR = os.path.join(SCRIPT_DIR, '【五六七普人口数据】上海区县层级的人口受教育情况')

# ====================================================================
# 五普: 28列, 表头行3-5 (1-indexed), 空行6, 数据从行7开始
# ====================================================================
def preprocess_edu_5():
    src = os.path.join(SRC_DIR, '【五普】人口受教育情况.xlsx')
    out = os.path.join(SCRIPT_DIR, 'education_5_cleaned.csv')

    COLUMNS = [
        '地区',
        '未上过学_男', '未上过学_女',
        '扫盲班_男', '扫盲班_女',
        '小学_男', '小学_女',
        '初中_男', '初中_女',
        '高中_男', '高中_女',
        '中专_男', '中专_女',
        '大学专科_男', '大学专科_女',
        '大学本科_男', '大学本科_女',
        '研究生_男', '研究生_女',
        '平均受教育年限_合计', '平均受教育年限_男', '平均受教育年限_女',
        '文盲人口_合计', '文盲人口_男', '文盲人口_女',
        '文盲率_合计', '文盲率_男', '文盲率_女',
    ]

    print("\n--- 五普受教育情况 ---")
    print(f"输入: {os.path.basename(src)}")

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb['Sheet1']

    records = []
    for row_idx in range(7, ws.max_row + 1):
        cell_a = ws.cell(row=row_idx, column=1).value
        cell_b = ws.cell(row=row_idx, column=2).value
        if cell_a is None or cell_b is None:
            continue
        row_data = []
        for col_idx in range(1, 29):  # 28列
            val = ws.cell(row=row_idx, column=col_idx).value
            row_data.append(val)
        records.append(row_data)

    df = pd.DataFrame(records, columns=COLUMNS)
    df['地区'] = df['地区'].astype(str).str.strip().str.replace(r'\s+', '', regex=True)

    num_cols = [c for c in COLUMNS if c != '地区']
    for col in num_cols:
        df[col] = pd.to_numeric(df[col], errors='coerce')

    # 行政层级
    def classify(name):
        if name in ('上海市',):
            return '全市合计'
        if name in ('市辖区', '市辖县'):
            return '市辖分类'
        if name.endswith('区') or name.endswith('县'):
            return '区县'
        return '其他'
    df['行政层级'] = df['地区'].apply(classify)
    df['普查'] = '五普'
    df['普查年份'] = 2000

    miss = df.isnull().sum()
    miss_any = miss[miss > 0]
    if len(miss_any) > 0:
        print(f"  ⚠ 缺失列: {dict(miss_any)}")
    else:
        print(f"  ✓ 无缺失值")

    df.to_csv(out, index=False, encoding='utf-8-sig')
    print(f"  输出: {os.path.basename(out)}, {len(df)} 行, {len(df.columns)} 列")
    print(f"  行政层级: {df['行政层级'].value_counts().to_dict()}")
    return df


# ====================================================================
# 六普: 22列, 表头行1-3 (1-indexed), 数据从行4开始
# ====================================================================
def preprocess_edu_6():
    src = os.path.join(SRC_DIR, '【六普】人口受教育情况.xlsx')
    out = os.path.join(SCRIPT_DIR, 'education_6_cleaned.csv')

    COLUMNS = [
        '地区',
        '未上过学_男', '未上过学_女',
        '小学_男', '小学_女',
        '初中_男', '初中_女',
        '高中_男', '高中_女',
        '大学专科_男', '大学专科_女',
        '大学本科及以上_男', '大学本科及以上_女',
        '平均受教育年限_合计', '平均受教育年限_男', '平均受教育年限_女',
        '文盲人口_合计', '文盲人口_男', '文盲人口_女',
        '文盲率_合计', '文盲率_男', '文盲率_女',
    ]

    print("\n--- 六普受教育情况 ---")
    print(f"输入: {os.path.basename(src)}")

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb['Sheet1']

    records = []
    for row_idx in range(4, ws.max_row + 1):
        cell_a = ws.cell(row=row_idx, column=1).value
        cell_b = ws.cell(row=row_idx, column=2).value
        if cell_a is None or cell_b is None:
            continue
        row_data = []
        for col_idx in range(1, 23):  # 22列
            val = ws.cell(row=row_idx, column=col_idx).value
            row_data.append(val)
        records.append(row_data)

    df = pd.DataFrame(records, columns=COLUMNS)
    df['地区'] = df['地区'].astype(str).str.strip().str.replace(r'\s+', '', regex=True)

    num_cols = [c for c in COLUMNS if c != '地区']
    for col in num_cols:
        df[col] = pd.to_numeric(df[col], errors='coerce')

    def classify(name):
        if name in ('上海市',):
            return '全市合计'
        if name in ('市辖区', '县'):
            return '市辖分类'
        if name.endswith('区') or name.endswith('县'):
            return '区县'
        return '其他'
    df['行政层级'] = df['地区'].apply(classify)
    df['普查'] = '六普'
    df['普查年份'] = 2010

    miss = df.isnull().sum()
    miss_any = miss[miss > 0]
    if len(miss_any) > 0:
        print(f"  ⚠ 缺失列: {dict(miss_any)}")
    else:
        print(f"  ✓ 无缺失值")

    df.to_csv(out, index=False, encoding='utf-8-sig')
    print(f"  输出: {os.path.basename(out)}, {len(df)} 行, {len(df.columns)} 列")
    print(f"  行政层级: {df['行政层级'].value_counts().to_dict()}")
    return df


# ====================================================================
# 七普: 24列, 表头行1-3 (1-indexed), 数据从行4开始
# ====================================================================
def preprocess_edu_7():
    src = os.path.join(SRC_DIR, '【七普】人口受教育情况.xlsx')
    out = os.path.join(SCRIPT_DIR, 'education_7_cleaned.csv')

    COLUMNS = [
        '地区',
        '未上过学_男', '未上过学_女',
        '学前教育_男', '学前教育_女',
        '小学_男', '小学_女',
        '初中_男', '初中_女',
        '高中_男', '高中_女',
        '大学专科_男', '大学专科_女',
        '大学本科及以上_男', '大学本科及以上_女',
        '平均受教育年限_合计', '平均受教育年限_男', '平均受教育年限_女',
        '文盲人口_合计', '文盲人口_男', '文盲人口_女',
        '文盲率_合计', '文盲率_男', '文盲率_女',
    ]

    print("\n--- 七普受教育情况 ---")
    print(f"输入: {os.path.basename(src)}")

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb['Sheet1']

    records = []
    for row_idx in range(4, ws.max_row + 1):
        cell_a = ws.cell(row=row_idx, column=1).value
        cell_b = ws.cell(row=row_idx, column=2).value
        if cell_a is None or cell_b is None:
            continue
        row_data = []
        for col_idx in range(1, 25):  # 24列
            val = ws.cell(row=row_idx, column=col_idx).value
            row_data.append(val)
        records.append(row_data)

    df = pd.DataFrame(records, columns=COLUMNS)
    df['地区'] = df['地区'].astype(str).str.strip().str.replace(r'\s+', '', regex=True)

    num_cols = [c for c in COLUMNS if c != '地区']
    for col in num_cols:
        df[col] = pd.to_numeric(df[col], errors='coerce')

    def classify(name):
        if name in ('上海市',):
            return '全市合计'
        if name in ('市辖区',):
            return '市辖分类'
        if name.endswith('区') or name.endswith('县'):
            return '区县'
        return '其他'
    df['行政层级'] = df['地区'].apply(classify)
    df['普查'] = '七普'
    df['普查年份'] = 2020

    miss = df.isnull().sum()
    miss_any = miss[miss > 0]
    if len(miss_any) > 0:
        print(f"  ⚠ 缺失列: {dict(miss_any)}")
    else:
        print(f"  ✓ 无缺失值")

    df.to_csv(out, index=False, encoding='utf-8-sig')
    print(f"  输出: {os.path.basename(out)}, {len(df)} 行, {len(df.columns)} 列")
    print(f"  行政层级: {df['行政层级'].value_counts().to_dict()}")
    return df


# ====================================================================
if __name__ == '__main__':
    print("=" * 60)
    print("  预处理: 区县层级人口受教育情况 (五/六/七普)")
    print("=" * 60)

    df5 = preprocess_edu_5()
    df6 = preprocess_edu_6()
    df7 = preprocess_edu_7()

    print(f"\n{'='*60}")
    print("  汇总")
    print(f"{'='*60}")
    print(f"  五普: {len(df5)} 行, {len(df5.columns)} 列")
    print(f"  六普: {len(df6)} 行, {len(df6.columns)} 列")
    print(f"  七普: {len(df7)} 行, {len(df7.columns)} 列")
    print("\n[全部完成]")
