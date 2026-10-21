# -*- coding: utf-8 -*-
"""WP13.11 - build the two standalone availability statements (v14).

v14 removes the "Data Availability Statement" heading and its paragraph from
the manuscript; the author wants the manuscript to contain paper content only.
The statement text is not rewritten: this module reads it out of
`docs/_v14_改写报告.json` (the exact string that left the manuscript) and splits
it into the two standalone documents

    manuscript\\Data_Availability_Statement.docx
    manuscript\\Code_Availability_Statement.docx

so the delivered wording is byte-identical to what was in v13.

Both files carry the manuscript title and corresponding author so that the
editorial office can match them to the submission.  The three placeholders
(Zenodo DOI, housing-price provider, GitHub repository) are preserved verbatim
until the author supplies the values.

Requires python-docx (bundled python), not the climada env.

USAGE
    <bundled python> tools\\make_declaration_docs.py
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

NEWROOT = Path(__file__).resolve().parent.parent
DOCS = NEWROOT / "docs"
MAN = NEWROOT / "manuscript"

REPORT = DOCS / "_v14_改写报告.json"
DA_OUT = MAN / "Data_Availability_Statement.docx"
CODE_OUT = MAN / "Code_Availability_Statement.docx"

TITLE = "Wind–surge asymmetry in typhoon risk to Shanghai's housing market"
AUTHOR = "Chen Jian, Fudan University International School of Finance"
JOURNAL = "npj Natural Hazards"


def build(path: Path, heading: str, label: str, body: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(heading, level=1)
    for k, v in (("Manuscript", TITLE), ("Corresponding author", AUTHOR),
                 ("Journal", JOURNAL)):
        p = doc.add_paragraph()
        r = p.add_run("%s: " % k)
        r.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(4)

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(body)
    doc.save(str(path))


def main() -> int:
    rep = json.loads(REPORT.read_text(encoding="utf-8"))
    raw = rep["removed_text"]["P253"]

    # split "Data availability. <data> Code availability. <code>"
    m = re.search(r"Code availability\.\s*", raw)
    if not m:
        raise SystemExit("could not find the Code availability label")
    data_part = raw[:m.start()].strip()
    code_part = raw[m.end():].strip()
    data_part = re.sub(r"^Data availability\.\s*", "", data_part)

    print("=== WP13.11 availability statements ===")
    print("  source paragraph: %d chars" % len(raw))
    print("  data part       : %d chars" % len(data_part))
    print("  code part       : %d chars" % len(code_part))

    # the split must reconstruct the source (modulo the two labels)
    rebuilt = ("Data availability. " + data_part + " Code availability. "
               + code_part)
    if rebuilt.replace("  ", " ") != raw.replace("  ", " "):
        print("  ! split does not reconstruct the source text")
        return 1

    build(DA_OUT, "Data Availability Statement", "Data availability.", data_part)
    build(CODE_OUT, "Code Availability Statement", "Code availability.", code_part)

    # numeric tokens carried out of the manuscript
    tok_re = re.compile(r"[+\u2212-]?\d+(?:\.\d+)?(?:[eE][+\u2212-]?\d+)?")
    n_data = len(tok_re.findall(data_part))
    n_code = len(tok_re.findall(code_part))
    print("  numeric tokens  : DA %d + code %d" % (n_data, n_code))
    for p in (DA_OUT, CODE_OUT):
        print("  wrote %s (%d B)" % (p.name, p.stat().st_size))

    summary = {
        "data_doc": {"name": DA_OUT.name, "bytes": DA_OUT.stat().st_size,
                     "tokens": n_data},
        "code_doc": {"name": CODE_OUT.name, "bytes": CODE_OUT.stat().st_size,
                     "tokens": n_code},
        "placeholders": re.findall(r"\[[^\]]+\]", raw),
    }
    (DOCS / "_v14_声明文件摘要.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8", newline="\n")
    print("RESULT: DECLARATIONS-BUILT")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
