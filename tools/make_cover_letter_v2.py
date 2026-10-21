#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""make_cover_letter_v2.py — WP8-3 · 由 FINAL 投稿信派生 v2，把 6 处数字/措辞对齐 v10。

为什么这样做
------------
`Cover_Letter_npjNaturalHazards_FINAL.docx`（17 段 / 495 词 / A4 / 1 英寸边距 / Times New Roman）
的数字是按**冻结底稿**写的；v10 换成官方层+岸线解耦（A 口径）后，有 6 处要改：

    P8  CNY 33.3 billion      -> CNY 33.4 billion      (33.3920，v10 P16 已印 33.4)
    P8  CNY 40.2 billion      -> CNY 40.3 billion      (40.2978)
    P8  undefended baseline   -> defended baseline     (39.5% = 26.5185/19.0056-1，对齐 v10 P16 "defended level")
    P9  (0.83‰)               -> (0.82‰)               (0.8214)
    P9  (0.68‰)               -> (0.67‰)               (0.6678)
    P9  of 0.38‰              -> of 0.39‰              (v10 P19 已印 0.39‰)

其余数字**一律不动**（24.2 / 19.0 / 26.5 / 63.3 / 21.5% / 74% / 26% / 40% / ≈16 bn / 39.5% /
585,033 / 24.87 M / 2.7 m / 500-path / 1995–2023 —— 逐项已在 `docs/_v10断言/_数值清单.csv` 复核）。

做法：**逐 run 文本替换**（python-docx），保留全部格式；绝不重建段落。
六个目标串都落在单 run 段（P8、P9 各 1 个 run），因此不存在跨 run 撕裂风险。

运行环境
--------
本脚本需要 **python-docx**，管道环境（climada_env）没有，必须用 bundled runtime：

    C:\\Users\\Jack\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe tools\\make_cover_letter_v2.py

退出码：0 成功 / 2 断言失败或环境缺失 / 3 源文件漂移
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Emu
except ImportError:  # pragma: no cover
    print(
        "ERROR: 本脚本需要 python-docx。请用 bundled runtime 运行：\n"
        r"  C:\Users\Jack\.dsh\dsh-runtimes\dsh-primary-runtime\dependencies\python\python.exe "
        r"tools\make_cover_letter_v2.py",
        file=sys.stderr,
    )
    raise SystemExit(2)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
JOURNAL_ROOT = PROJECT_ROOT.parent                      # …\Nature-JF投稿
SRC = JOURNAL_ROOT / "Nature Cities" / "提交审稿人" / "Cover_Letter_npjNaturalHazards_FINAL.docx"
DST = PROJECT_ROOT / "manuscript" / "Cover_Letter_npjNaturalHazards_v2.docx"
REPORT = PROJECT_ROOT / "docs" / "_cover_letter_v2_check.json"

# ── 期望结构（源件实测） ──────────────────────────────────────────────────
EXPECT_PARAS = 17
EXPECT_WORDS = 495
EXPECT_SECTIONS = 1
A4_W, A4_H = 7562215, 10689590          # EMU（= 8.27in x 11.69in）
MARGIN_EMU = 914400                     # 1 英寸

# ── 6 处编辑：(old, new, 依据说明) ─────────────────────────────────────────
EDITS = [
    ("CNY 33.3 billion", "CNY 33.4 billion", "33.3920 -> 33.4（v10 P16）"),
    ("CNY 40.2 billion", "CNY 40.3 billion", "40.2978 -> 40.3（v10 P16）"),
    ("undefended baseline", "defended baseline", "39.5% = 26.5185/19.0056-1，对齐 v10 P16 'defended level'"),
    ("(0.83‰)", "(0.82‰)", "0.8214 -> 0.82（v10 P19）"),
    ("(0.68‰)", "(0.67‰)", "0.6678 -> 0.67（v10 P19）"),
    ("of 0.38‰", "of 0.39‰", "v10 P19 已印 0.39‰"),
]

# ── 不该出现的旧串 ───────────────────────────────────────────────────────
FORBIDDEN_AFTER = ["CNY 33.3 billion", "CNY 40.2 billion", "undefended baseline",
                   "0.83‰", "0.68‰", "0.38‰"]
# ── 必须仍在的不变量串（证明没顺手改坏别的） ────────────────────────────────
MUST_KEEP = ["CNY 24.2 billion", "CNY 19.0 billion", "CNY 26.5 billion", "CNY 63.3 trillion",
             "21.5%", "74%", "26%", "40%", "39.5%", "585,033", "24.87 million", "2.7 m",
             "500-path", "1995–2023", "0.77‰", "nearly CNY 16 billion"]
# ── 新串必须在 ────────────────────────────────────────────────────────────
MUST_APPEAR = [e[1] for e in EDITS]


def die(msg: str, code: int = 2) -> "NoReturn":  # noqa: F821
    print("ERROR: " + msg, file=sys.stderr)
    raise SystemExit(code)


def para_words(doc) -> int:
    return sum(len(p.text.split()) for p in doc.paragraphs)


def structural(doc) -> dict:
    sec = doc.sections[0]
    return {
        "paras": len(doc.paragraphs),
        "words": para_words(doc),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "page_w_emu": int(sec.page_width),
        "page_h_emu": int(sec.page_height),
        "margin_left_emu": int(sec.left_margin),
        "margin_right_emu": int(sec.right_margin),
    }


def apply_edits(doc) -> list:
    """逐 run 替换。返回命中记录 [(para_idx, old, new, count)]。"""
    log = []
    for old, new, why in EDITS:
        hits = 0
        for pi, par in enumerate(doc.paragraphs):
            for run in par.runs:
                if old in run.text:
                    n = run.text.count(old)
                    run.text = run.text.replace(old, new)
                    hits += n
                    log.append({"para": pi, "old": old, "new": new, "count": n, "why": why})
        if hits != 1:
            die("编辑 %r 命中 %d 次（应为 1 次）——源件可能已漂移" % (old, hits))
    return log


def assertions(doc) -> list:
    checks = []
    text = "\n".join(p.text for p in doc.paragraphs)
    st = structural(doc)

    def chk(name, ok, detail=""):
        checks.append({"check": name, "pass": bool(ok), "detail": detail})
        return bool(ok)

    chk("段数 = %d" % EXPECT_PARAS, st["paras"] == EXPECT_PARAS, str(st["paras"]))
    chk("词数 = %d" % EXPECT_WORDS, st["words"] == EXPECT_WORDS, str(st["words"]))
    chk("section 数 = 1", st["sections"] == EXPECT_SECTIONS, str(st["sections"]))
    chk("A4 页面", (st["page_w_emu"], st["page_h_emu"]) == (A4_W, A4_H),
        "%dx%d" % (st["page_w_emu"], st["page_h_emu"]))
    chk("左右边距 1 英寸",
        st["margin_left_emu"] == MARGIN_EMU and st["margin_right_emu"] == MARGIN_EMU,
        "%d/%d" % (st["margin_left_emu"], st["margin_right_emu"]))
    for s in FORBIDDEN_AFTER:
        chk("旧串已消失 %r" % s, s not in text)
    for s in MUST_KEEP:
        chk("不变量仍在 %r" % s, s in text)
    for s in MUST_APPEAR:
        chk("新串已在 %r" % s, s in text)
    # 占位符检查：只允许日期段那种【】包裹
    import re
    placeholders = re.findall(r"\[[^\]]{3,60}\]", text)
    chk("无 [] 残留占位符", not placeholders, "; ".join(placeholders[:3]))
    return checks


def main() -> int:
    ap = argparse.ArgumentParser(description="派生 Cover_Letter v2（数字对齐 v10）")
    ap.add_argument("--src", type=Path, default=SRC)
    ap.add_argument("--dst", type=Path, default=DST)
    ap.add_argument("--check", action="store_true", help="只校验已有 dst，不重写")
    args = ap.parse_args()

    if args.check:
        if not args.dst.is_file():
            die("找不到 %s，请先不带 --check 运行一次" % args.dst)
        doc = Document(str(args.dst))
        checks = assertions(doc)
        print("== 校验 %s ==" % args.dst.name)
        bad = 0
        for c in checks:
            print("  [%s] %s%s" % ("PASS" if c["pass"] else "FAIL", c["check"],
                                   ("  (" + c["detail"] + ")") if c["detail"] else ""))
            bad += 0 if c["pass"] else 1
        print("\n%s  (%d PASS / %d FAIL)" % ("RESULT: PASS" if not bad else "RESULT: FAIL",
                                             len(checks) - bad, bad))
        return 0 if not bad else 2

    if not args.src.is_file():
        die("找不到源件：%s" % args.src)

    doc = Document(str(args.src))
    src_st = structural(doc)
    print("== 源件 ==")
    print("  %s (%d B)" % (args.src.name, args.src.stat().st_size))
    print("  paras=%d words=%d sections=%d page=%dx%d margins=%d/%d"
          % (src_st["paras"], src_st["words"], src_st["sections"],
             src_st["page_w_emu"], src_st["page_h_emu"],
             src_st["margin_left_emu"], src_st["margin_right_emu"]))
    if src_st["paras"] != EXPECT_PARAS or src_st["words"] != EXPECT_WORDS:
        die("源件结构与期望不符（paras=%d words=%d，期望 %d/%d）——请确认是否换了源件"
            % (src_st["paras"], src_st["words"], EXPECT_PARAS, EXPECT_WORDS), 3)

    log = apply_edits(doc)
    print("\n== 6 处编辑（逐 run 替换） ==")
    for r in log:
        print("  P%-2d %-22r -> %-22r  x%d   # %s"
              % (r["para"], r["old"], r["new"], r["count"], r["why"]))

    checks = assertions(doc)
    print("\n== 断言 ==")
    bad = 0
    for c in checks:
        print("  [%s] %s%s" % ("PASS" if c["pass"] else "FAIL", c["check"],
                               ("  (" + c["detail"] + ")") if c["detail"] else ""))
        bad += 0 if c["pass"] else 1
    if bad:
        die("断言失败 %d 项，未写出文件" % bad)

    args.dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(args.dst))
    print("\n已写出: %s (%d B)" % (args.dst, args.dst.stat().st_size))

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "src": str(args.src),
        "src_bytes": args.src.stat().st_size,
        "dst": str(args.dst),
        "dst_bytes": args.dst.stat().st_size,
        "edits": log,
        "src_structure": src_st,
        "dst_structure": structural(doc),
        "checks": checks,
        "checks_pass": sum(1 for c in checks if c["pass"]),
        "checks_fail": bad,
    }
    REPORT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("校验记录: %s" % REPORT)
    print("\nRESULT: PASS  (%d PASS / 0 FAIL)" % len(checks))
    return 0


if __name__ == "__main__":
    sys.exit(main())
